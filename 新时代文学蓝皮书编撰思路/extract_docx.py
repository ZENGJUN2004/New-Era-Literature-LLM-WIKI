import docx
import sys, os

base = r"D:\zjun\新时代文学LLMWIKI\新时代文学蓝皮书编撰思路"
files = [
    "蓝皮书专家优化意见稿.docx",
    "260906下午蓝皮书内部研讨会录音整理稿（AI整理后初修版） 汪静波.docx",
    "260906 蓝皮书内部研讨会纪要 汪静波(1).docx",
]

out_dir = os.path.join(base, "_extract")
os.makedirs(out_dir, exist_ok=True)

for f in files:
    path = os.path.join(base, f)
    doc = docx.Document(path)
    parts = []
    # paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            style = p.style.name if p.style else ""
            parts.append(f"[{style}] {txt}")
    # tables
    for ti, t in enumerate(doc.tables):
        parts.append(f"\n=== 表格 {ti+1} ===")
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    out_name = f.replace(".docx", "") + ".txt"
    with open(os.path.join(out_dir, out_name), "w", encoding="utf-8") as fh:
        fh.write("\n".join(parts))
    print(f"提取完成: {out_name}  段落/行数约: {len(parts)}")
